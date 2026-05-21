import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from derivacion_drm.api.app import create_app

OUT = Path(__file__).resolve().parents[2] / "smoke_output.docx"

client = TestClient(create_app())

payload = {
    "caso": {
        "adolescente": {
            "nombre": "Juan Pérez García",
            "run": "22.846.782-0",
            "genero": "Masculino",
            "domicilio": "Calle Falsa 123",
            "comuna": "Lo Espejo",
        },
        "causa": {
            "tribunal": "Unidad Especializada RPA",
            "tipo_resolucion": "oficio",
            "ruc": "2600664806-0",
            "rit": "2903-2026",
            "delito": "robo con violencia",
            "fecha_resolucion": "2026-05-10",
            "art_37_bis": True,
        },
        "medida": "MCA",
        "adulto": {"nombre": "Carmen Soto", "telefono": "+56999999999"},
        "centro": {
            "nombre": "MCA Centro Sur",
            "director": "Pedro Soto",
            "mail": "pedro@centro.cl",
            "telefono": "22 111 2222",
            "direccion": "Av Sur 100",
            "tipo": "Priorizada",
        },
    },
    "profesional": "Juan Manuel Olivares Oyarzún",
    "fecha_emision": "2026-05-20",
}

r = client.post("/api/derivacion/generar", json=payload)
assert r.status_code == 200, f"generar fallo: {r.status_code} {r.text}"
OUT.write_bytes(r.content)
print(f"WROTE {OUT}")

# Verify the Word by reopening it
d = Document(str(OUT))
all_text = "\n".join(p.text for p in d.paragraphs)
assert "JUAN P" in all_text.upper(), "nombre no aparece en mayúsculas"
assert "Lo Espejo" in all_text, "comuna ausente"
assert "Carmen Soto" in all_text, "adulto responsable ausente"
assert "artículo 37 bis" in all_text.lower() or "37 bis" in all_text.lower(), "37 bis ausente"
assert "MCA Centro Sur" in all_text, "centro ausente"
# Footer
footer_texts = [p.text for sec in d.sections for p in sec.footer.paragraphs]
assert any("Pedro De Valdivia" in t for t in footer_texts), "pie de pagina ausente"
# Bookman font on some run
fonts = {run.font.name for p in d.paragraphs for run in p.runs if run.font.name}
assert "Bookman Old Style" in fonts, "fuente ausente"
# Margenes 2.5cm
sec = d.sections[0]
for which in ("top", "bottom", "left", "right"):
    val = getattr(sec, f"{which}_margin").cm
    assert abs(val - 2.5) < 0.05, f"margen {which} no es 2.5cm: {val}"

# Filename header
cd_header = r.headers.get("content-disposition", "")
assert "Juan_Perez_Garcia_MCA_Unidad_Especializada_Rpa_2903-2026_2600664806-0.docx" in cd_header, \
    f"filename incorrecto: {cd_header}"

print("ALL ASSERTIONS PASSED")
