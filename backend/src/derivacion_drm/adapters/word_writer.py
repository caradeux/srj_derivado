import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from derivacion_drm.domain.word_spec import (
    ParagraphSpec, RunSpec, TableSpec, WordDocumentSpec,
)

FONT_NAME = "Bookman Old Style"
FONT_SIZE_PT = 12

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _aplicar_font(run, *, bold: bool, underline: bool):
    run.font.name = FONT_NAME
    # Asegurar font para Asia / complex scripts también
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT_NAME)
    run.font.size = Pt(FONT_SIZE_PT)
    run.bold = bold
    run.underline = underline


def _add_hyperlink(paragraph, url: str, run_spec: RunSpec):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    # Estilo: azul + negrita + subrayado
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000EE")
    rPr.append(color)
    b = OxmlElement("w:b")
    rPr.append(b)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi"):
        rFonts.set(qn(attr), FONT_NAME)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(FONT_SIZE_PT * 2))
    rPr.append(sz)
    new_run.append(rPr)
    text = OxmlElement("w:t")
    text.text = run_spec.text.upper() if run_spec.uppercase else run_spec.text
    text.set(qn("xml:space"), "preserve")
    new_run.append(text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _aplicar_parrafo(doc, spec: ParagraphSpec):
    p = doc.add_paragraph()
    p.alignment = _ALIGN[spec.alignment]
    pf = p.paragraph_format
    pf.line_spacing = spec.line_spacing
    if spec.indent_first_line_cm is not None:
        pf.first_line_indent = Cm(spec.indent_first_line_cm)
    for r in spec.runs:
        if r.hyperlink:
            _add_hyperlink(p, r.hyperlink, r)
            continue
        texto = r.text.upper() if r.uppercase else r.text
        run = p.add_run(texto)
        _aplicar_font(run, bold=r.bold, underline=r.underline)


def _aplicar_tabla(doc, spec: TableSpec):
    table = doc.add_table(rows=len(spec.rows), cols=2)
    table.style = spec.style
    for i, (etiqueta, valor) in enumerate(spec.rows):
        for j, txt in enumerate((etiqueta, valor)):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            _aplicar_font(run, bold=(j == 0), underline=False)


def _setup_documento(doc):
    """Márgenes 2.5cm, default font Bookman 12pt."""
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)


def _agregar_footer(doc, texto: str):
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        _aplicar_font(run, bold=False, underline=False)


def _agregar_logo(doc, logo_path: Path | None):
    if logo_path is None or not logo_path.exists():
        return
    sec = doc.sections[0]
    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(logo_path), width=Cm(4))


def escribir_word(spec: WordDocumentSpec, logo_path: Path | None) -> bytes:
    """Aplica formato institucional al WordDocumentSpec y devuelve bytes .docx."""
    doc = Document()
    _setup_documento(doc)
    if spec.incluir_logo:
        _agregar_logo(doc, logo_path)
    for item in spec.paragraphs:
        if isinstance(item, ParagraphSpec):
            _aplicar_parrafo(doc, item)
        elif isinstance(item, TableSpec):
            _aplicar_tabla(doc, item)
    _agregar_footer(doc, spec.footer_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
