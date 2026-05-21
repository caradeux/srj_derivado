from pathlib import Path

from docx import Document

from derivacion_drm.domain.models import DocumentoCargado, Tabla


def leer_docx(path: Path) -> DocumentoCargado:
    if not path.exists():
        raise FileNotFoundError(path)
    d = Document(str(path))
    texto = "\n".join(p.text for p in d.paragraphs)
    tablas: list[Tabla] = []
    for t in d.tables:
        if not t.rows:
            continue
        encab = [c.text.strip() for c in t.rows[0].cells]
        filas = [[c.text.strip() for c in fila.cells] for fila in t.rows[1:]]
        tablas.append(Tabla(encabezados=encab, filas=filas))
    return DocumentoCargado(nombre_archivo=path.name, texto=texto, tablas=tablas)
