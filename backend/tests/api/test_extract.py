import io

from fastapi.testclient import TestClient

from derivacion_drm.api.app import create_app


def _docx_minimo_bytes() -> bytes:
    from docx import Document
    d = Document()
    d.add_paragraph("RUC 2600664806-0 RIT 2903-2026")
    tabla = d.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "NOMBRE IMPUTADO"
    tabla.cell(0, 1).text = "RUT"
    tabla.cell(1, 0).text = "Juan Pérez"
    tabla.cell(1, 1).text = "22.846.782-0"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_docx_devuelve_extraction_result():
    client = TestClient(create_app())
    resp = client.post(
        "/api/extract",
        files=[("files", ("acta.docx", _docx_minimo_bytes(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["causa"]["ruc"] == "2600664806-0"
    assert data["causa"]["rit"] == "2903-2026"
    assert len(data["candidatos"]) == 1
    assert data["candidatos"][0]["nombre"] == "Juan Pérez"


def test_extract_sin_archivos_400():
    client = TestClient(create_app())
    resp = client.post("/api/extract", files=[])
    assert resp.status_code in (400, 422)


def test_extract_archivo_no_soportado():
    client = TestClient(create_app())
    resp = client.post(
        "/api/extract",
        files=[("files", ("nota.txt", b"x", "text/plain"))],
    )
    assert resp.status_code == 400
