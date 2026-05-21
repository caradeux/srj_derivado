from pathlib import Path
import pytest

FIXTURES = Path(__file__).parent / "fixtures"

@pytest.fixture
def fixtures_path() -> Path:
    return FIXTURES


@pytest.fixture
def sample_pdf_path(tmp_path: Path) -> Path:
    """Genera un PDF mínimo con texto reconocible para los extractors."""
    try:
        from reportlab.pdfgen import canvas
    except ImportError:
        pytest.skip("reportlab no instalado; skipping PDF fixture test")
    p = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(p))
    c.drawString(100, 750, "RUC 2600664806-0 RIT 2903-2026")
    c.drawString(100, 730, "Juan Perez 22.846.782-0 Lo Espejo")
    c.save()
    return p


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    from docx import Document
    p = tmp_path / "sample.docx"
    d = Document()
    d.add_paragraph("RUC 2600664806-0 RIT 2903-2026")
    tabla = d.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "NOMBRE IMPUTADO"
    tabla.cell(0, 1).text = "RUT"
    tabla.cell(1, 0).text = "Juan Pérez"
    tabla.cell(1, 1).text = "22.846.782-0"
    d.save(str(p))
    return p


@pytest.fixture
def sample_excel_path(tmp_path: Path) -> Path:
    import pandas as pd
    p = tmp_path / "catalogo.xlsx"
    df = pd.DataFrame({
        "PROGRAMA": ["MCA Centro Sur", "LAE Oriente"],
        "Tipo de medida o sanción": ["MCA", "LAE"],
        "COMUNAS PRIORIZADAS": ["Lo Espejo, La Cisterna", "Peñalolén"],
        "COMUNAS NO PRIORIZADAS": ["", "La Reina, Las Condes"],
        "DIRECCION": ["Av Sur 100", "Av Oriente 200"],
        "DIRECTOR": ["Pedro Soto", "Ana Vidal"],
        "CONTACTO MAIL": ["pedro@x.cl", "ana@y.cl"],
        "TELEFONO": ["22 111", "22 222"],
    })
    with pd.ExcelWriter(p) as w:
        df.to_excel(w, sheet_name="Desacumulado", index=False)
    return p
