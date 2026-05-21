import io
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from derivacion_drm.adapters.word_writer import escribir_word
from derivacion_drm.domain.word_spec import (
    ParagraphSpec, RunSpec, WordDocumentSpec,
)


def _spec_minimo() -> WordDocumentSpec:
    return WordDocumentSpec(
        paragraphs=[
            ParagraphSpec(
                runs=[RunSpec(text="Hola "), RunSpec(text="MUNDO", bold=True, uppercase=True)],
                alignment="center", indent_first_line_cm=1.25,
            ),
        ],
        incluir_logo=False,
        filename="x.docx",
    )


def test_escribir_word_devuelve_bytes_no_vacios():
    data = escribir_word(_spec_minimo(), logo_path=None)
    assert isinstance(data, bytes) and len(data) > 1000


def test_word_aplica_font_bookman_y_margenes(tmp_path: Path):
    data = escribir_word(_spec_minimo(), logo_path=None)
    out = tmp_path / "out.docx"
    out.write_bytes(data)
    d = Document(str(out))
    # Márgenes 2.5 cm en todos los bordes
    sec = d.sections[0]
    assert abs(sec.top_margin.cm - 2.5) < 0.05
    assert abs(sec.left_margin.cm - 2.5) < 0.05
    assert abs(sec.right_margin.cm - 2.5) < 0.05
    assert abs(sec.bottom_margin.cm - 2.5) < 0.05
    # Font Bookman Old Style 12pt en al menos un run
    fonts = {r.font.name for p in d.paragraphs for r in p.runs if r.font.name}
    assert "Bookman Old Style" in fonts


def test_word_pie_de_pagina_contiene_direccion(tmp_path: Path):
    data = escribir_word(_spec_minimo(), logo_path=None)
    out = tmp_path / "out.docx"
    out.write_bytes(data)
    d = Document(str(out))
    footer_texts = [
        p.text for sec in d.sections for p in sec.footer.paragraphs
    ]
    assert any("Pedro De Valdivia" in t for t in footer_texts)


def test_word_aplica_uppercase_a_runs_marcados(tmp_path: Path):
    data = escribir_word(_spec_minimo(), logo_path=None)
    out = tmp_path / "out.docx"
    out.write_bytes(data)
    d = Document(str(out))
    todos = "".join(p.text for p in d.paragraphs)
    assert "MUNDO" in todos
